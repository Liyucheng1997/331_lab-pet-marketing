"""Generate the two operating workbooks in the style of the user's own 商品进价表 / 销售统计表.

    python tools/build_workbooks.py [--price-list 01_商品进价表.xlsx] [--sales 02_销售统计表.xlsx]

1. 选品调研对比表.xlsx      — one row per candidate: three-platform price/sales, unit profit, monthly estimate.
2. 上新排期与销售统计表.xlsx — 上新排期 (products + listing status), 销售统计 (orders), 月度复盘 (summary).

If the user's existing workbooks are found they are copied in, so the new files start with real data.
Output goes to excel/ (git-ignored because it contains business data).
"""
from pathlib import Path
import argparse, datetime as dt
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent.parent / 'excel'
SOURCE = Path('F:/电商工作流')
STORE = 'F1 CASA'
EUR = '[$€-2]\\ #,##0.00'
PCT = '0%'
DATE = 'yyyy\\-mm\\-dd;@'
MONTH = 'yyyy\\-mm'
TEXT = '@'

# Same look as the user's sheets: Calibri, red 20pt title, light gold header with medium border, thin grid, gold total row.
TITLE_FONT = Font(name='Calibri', size=20, color='FFFF0000')
FONT = Font(name='Calibri', size=11)
HEAD_FILL = PatternFill('solid', fgColor=Color(theme=7, tint=0.5999938962981048))
TOTAL_FILL = PatternFill('solid', fgColor='FFFFC000')
GREEN = PatternFill('solid', fgColor='FF00B050', bgColor='FF00B050')
YELLOW = PatternFill('solid', fgColor='FFFFFF00', bgColor='FFFFFF00')
RED = PatternFill('solid', fgColor='FFFF7C80', bgColor='FFFF7C80')
MED, THIN = Side(style='medium'), Side(style='thin')
HEAD_BORDER = Border(left=MED, right=MED, top=MED, bottom=MED)
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)


def sheet(ws, title, columns, first_data_row=3):
    """columns: list of (header, width, number_format)."""
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
    ws['A1'] = title; ws['A1'].font = TITLE_FONT; ws['A1'].alignment = CENTER
    ws.row_dimensions[1].height = 27.75
    ws.row_dimensions[2].height = 33
    header(ws, 2, [c[0] for c in columns])
    for i, (_, width, _) in enumerate(columns, 1): ws.column_dimensions[L(i)].width = width
    ws.freeze_panes = ws.cell(first_data_row, 1)


def header(ws, r, labels):
    for i, head in enumerate(labels, 1):
        c = ws.cell(r, i, head); c.font = FONT; c.fill = HEAD_FILL; c.border = HEAD_BORDER; c.alignment = CENTER


def row(ws, r, values, columns, fill=None):
    ws.row_dimensions[r].height = 20.1
    for i, ((_, _, fmt), v) in enumerate(zip(columns, values), 1):
        c = ws.cell(r, i, v); c.font = FONT; c.border = CELL_BORDER; c.alignment = CENTER
        if fmt: c.number_format = fmt
        if fill: c.fill = fill


def dropdown(ws, options, ref):
    dv = DataValidation(type='list', formula1='"' + ','.join(options) + '"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(ref)


def color_words(ws, ref, words):
    for word, fill in words.items():
        ws.conditional_formatting.add(ref, CellIsRule(operator='equal', formula=[f'"{word}"'], fill=fill))


def read_rows(path, first=3):
    """Data rows of one of the user's sheets (stops at the 总计 row or the first empty row)."""
    if not path or not Path(path).is_file(): return []
    ws = load_workbook(path).active
    out = []
    for r in range(first, ws.max_row + 1):
        values = [ws.cell(r, c).value for c in range(1, 16)]
        if values[1] in (None, '') or values[1] == '总计': break
        out.append(values)
    return out


# ---------------------------------------------------------------- 1. 选品调研对比表
RESEARCH = [('目录', 6.7, None), ('调研日期', 12, DATE), ('中文简介', 26, None), ('意大利语关键词', 26, None),
            ('税前进货价', 11.7, EUR), ('税后进货价', 11.7, EUR),
            ('AliExpress售价', 12, EUR), ('AliExpress月销量', 12, '0'), ('Amazon售价', 12, EUR), ('Amazon月销量', 12, '0'),
            ('TikTok售价', 12, EUR), ('TikTok月销量', 12, '0'), ('三平台最低价', 12, EUR), ('三平台月销量合计', 12, '0'),
            ('我方售价', 11.7, EUR), ('买家物流费', 11.7, EUR), ('物流费用', 11.7, EUR), ('平台佣金8%+物流佣金5%', 16, EUR),
            ('IVA报税', 11.7, EUR), ('单件利润', 11.7, EUR), ('利润率', 9, PCT), ('预估月销量', 11.7, '0'), ('预估月利润', 12, EUR),
            ('结论', 10, None), ('备注', 30, None)]


def build_research(rows=60):
    wb = Workbook(); ws = wb.active; ws.title = '选品调研'
    sheet(ws, f'{STORE} 选品调研对比表', RESEARCH)
    example = [1, dt.date(2026, 9, 20), 'LIFEman 7层省空间鞋架', 'scarpiera salvaspazio', 5.54, None,
               16.99, 320, 24.99, 150, 18.9, 60, None, None, 14.99, 3.89, 4.1, None, None, None, None, 20, None, '上新',
               '示例：进价和售价来自你的进价表，竞品数据是示例，可删除']
    for i in range(rows):
        r = 3 + i
        v = list(example) if i == 0 else [i + 1] + [None] * (len(RESEARCH) - 1)
        v[5] = f'=IF(E{r}="","",E{r}*1.22)'
        v[12] = f'=IF(COUNT(G{r},I{r},K{r})=0,"",MIN(G{r},I{r},K{r}))'
        v[13] = f'=IF(COUNT(H{r},J{r},L{r})=0,"",SUM(H{r},J{r},L{r}))'
        v[17] = f'=IF(O{r}="","",0.08*O{r}+0.05*P{r})'
        v[18] = f'=IF(O{r}="","",O{r}*20%)'
        v[19] = f'=IF(O{r}="","",O{r}+P{r}-F{r}-Q{r}-R{r}-S{r})'
        v[20] = f'=IF(O{r}="","",T{r}/O{r})'
        v[22] = f'=IF(OR(T{r}="",V{r}=""),"",T{r}*V{r})'
        row(ws, r, v, RESEARCH)
    last = 2 + rows
    dropdown(ws, ['上新', '观察', '放弃'], f'X3:X{last}')
    color_words(ws, f'X3:X{last}', {'上新': GREEN, '观察': YELLOW, '放弃': RED})
    ws.conditional_formatting.add(f'T3:T{last}', CellIsRule(operator='lessThan', formula=['0'], fill=RED))
    return wb


# ---------------------------------------------------------------- 2. 上新排期与销售统计表
PLAN = [('目录', 6.7, None), ('计划上架日期', 12, DATE), ('商品EAN码', 16.4, '0'), ('商品ID', 20, TEXT), ('商品简介', 44, None),
        ('中文简介', 28, None), ('税前进货价', 11.7, EUR), ('税后进货价', 11.7, EUR), ('电商原价', 11.7, EUR), ('日常促销价', 11.7, EUR),
        ('仓库位置', 11.7, None), ('仓库库存', 11.7, '#,##0'), ('作图', 9, None), ('AliExpress', 11, None), ('Amazon', 11, None),
        ('TikTok', 11, None), ('备注', 24, None)]
SALES = [('目录', 6.7, None), ('时间', 12, DATE), ('平台', 11, None), ('SKU编码', 20, None), ('销售价格', 11.7, EUR), ('买家物流费', 11.7, EUR),
         ('税后进货价格', 12.5, EUR), ('物流费用', 11.7, EUR), ('平台佣金', 11.7, EUR), ('联盟佣金', 11.7, EUR), ('IVA报税', 11.7, EUR),
         ('最终利润', 11.7, EUR), ('利润率', 9, PCT), ('异常情况', 24, None)]
STATUS = ['未开始', '进行中', '完成', '下架']


def build_ops(price_list=None, sales_file=None, plan_rows=150, sales_rows=500):
    wb = Workbook()

    # 上新排期 — the user's 商品进价表 plus a planned date, image status and one status column per platform.
    ws = wb.active; ws.title = '上新排期'
    sheet(ws, f'{STORE} 上新排期表', PLAN)
    source = read_rows(price_list)
    for i in range(plan_rows):
        r = 3 + i
        if i < len(source):
            s = source[i]  # 目录, EAN, 商品ID, 商品简介, 中文简介, 税前, 税后, 百货价, 销售方式, 原价, 促销价, 位置, 库存, 销售状态, 上品状态
            done = s[14] == '完成'
            ae = '下架' if s[13] == '下架' else '完成' if done else None
            v = [i + 1, None, s[1], str(s[2] or ''), s[3], s[4], s[5], None, s[9], s[10], s[11], s[12], '完成' if done else None, ae, None, None, None]
        elif i == 0:
            v = [1, dt.date(2026, 10, 1), 8388776542996, '1005011833238348', 'Scarpiera Salvaspazio LIFEman a 7 Ripiani', 'LIFEman 7层省空间鞋架',
                 5.54, None, 19.99, 14.99, 'Reggio', 640, '完成', '完成', '进行中', '未开始', '示例，可删除']
        else:
            v = [i + 1] + [None] * (len(PLAN) - 1)
        v[7] = f'=IF(G{r}="","",G{r}*1.22)'
        row(ws, r, v, PLAN)
    last = 2 + plan_rows
    for col in 'MNOP':
        dropdown(ws, STATUS, f'{col}3:{col}{last}')
        color_words(ws, f'{col}3:{col}{last}', {'完成': GREEN, '进行中': YELLOW, '下架': RED})

    # 销售统计 — the user's own columns and formulas, plus 平台. Total row sits right under the header.
    ss = wb.create_sheet('销售统计')
    sheet(ss, f'{STORE} 销售统计表', SALES, first_data_row=4)
    first, last_s = 4, 3 + sales_rows
    total = ['', '总计', '', ''] + [f'=SUM({c}{first}:{c}{last_s})' for c in 'EFGHIJKL'] + ['=IF(E3=0,"",L3/E3)', '']
    row(ss, 3, total, SALES, fill=TOTAL_FILL)
    orders = read_rows(sales_file)
    for i in range(sales_rows):
        r = first + i
        if i < len(orders):
            o = orders[i]  # 目录, 时间, 商品图片, SKU, 售价, 买家物流费, 进货, 物流, 佣金, 联盟, IVA, 利润, 利润率, 异常
            sku = str(o[3]) if isinstance(o[3], int) and len(str(o[3])) > 15 else o[3]  # 16-digit IDs exceed Excel precision
            v = [i + 1, o[1], 'AliExpress', sku, o[4], o[5], o[6], o[7], None, o[9], None, None, None, o[13]]
        elif i == 0:
            v = [1, dt.date(2026, 9, 21), 'AliExpress', 8388776542996, 14.99, 3.89, 6.76, 4.1, None, None, None, None, None, '示例，可删除']
        else:
            v = [i + 1] + [None] * (len(SALES) - 1)
        v[8] = f'=IF(E{r}="","",IF(C{r}="Amazon",0.15*(E{r}+F{r}),IF(C{r}="TikTok",0.09*(E{r}+F{r}),0.08*E{r}+0.05*F{r})))'
        v[10] = f'=IF(E{r}="","",E{r}*20%)'
        v[11] = f'=IF(E{r}="","",E{r}+F{r}-G{r}-H{r}-I{r}-J{r}-K{r})'
        v[12] = f'=IF(E{r}="","",L{r}/E{r})'
        row(ss, r, v, SALES)
    dropdown(ss, ['AliExpress', 'Amazon', 'TikTok'], f'C{first}:C{last_s}')
    ss.conditional_formatting.add(f'L{first}:L{last_s}', CellIsRule(operator='lessThan', formula=['0'], fill=RED))

    # 月度复盘 — two small summary tables and one chart, all read from 销售统计.
    rv = wb.create_sheet('月度复盘')
    MONTHLY = [('月份', 12, MONTH), ('订单数', 11.7, '0'), ('销售额', 12, EUR), ('最终利润', 12, EUR), ('利润率', 9, PCT)]
    sheet(rv, f'{STORE} 月度复盘', MONTHLY)
    dates = [o[1] for o in orders if isinstance(o[1], (dt.date, dt.datetime))]
    start = min(dates) if dates else dt.date(2026, 5, 1)
    S = "'销售统计'!"
    for i in range(12):
        r = 3 + i
        month = dt.date(start.year, start.month, 1) if i == 0 else f'=EDATE(A{r - 1},1)'
        rng = f'{S}$B:$B,">="&A{r},{S}$B:$B,"<"&EDATE(A{r},1)'
        row(rv, r, [month, f'=COUNTIFS({rng})', f'=SUMIFS({S}$E:$E,{rng})', f'=SUMIFS({S}$L:$L,{rng})', f'=IF(C{r}=0,"",D{r}/C{r})'], MONTHLY)
    row(rv, 15, ['总计', '=SUM(B3:B14)', '=SUM(C3:C14)', '=SUM(D3:D14)', '=IF(C15=0,"",D15/C15)'], MONTHLY, fill=TOTAL_FILL)
    header(rv, 17, ['平台', '订单数', '销售额', '最终利润', '利润率'])
    for i, p in enumerate(['AliExpress', 'Amazon', 'TikTok']):
        r = 18 + i
        row(rv, r, [p, f'=COUNTIF({S}$C:$C,A{r})', f'=SUMIFS({S}$E:$E,{S}$C:$C,A{r})', f'=SUMIFS({S}$L:$L,{S}$C:$C,A{r})', f'=IF(C{r}=0,"",D{r}/C{r})'],
            [('', 0, None)] + MONTHLY[1:])
    chart = BarChart(); chart.title = '每月销售额与利润'; chart.height, chart.width = 8, 16
    chart.add_data(Reference(rv, min_col=3, max_col=4, min_row=2, max_row=14), titles_from_data=True)
    chart.set_categories(Reference(rv, min_col=1, min_row=3, max_row=14))
    rv.add_chart(chart, 'G2')
    wb.active = 0
    return wb


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--price-list', default=str(SOURCE / '01_商品进价表.xlsx'))
    ap.add_argument('--sales', default=str(SOURCE / '02_销售统计表.xlsx'))
    a = ap.parse_args()
    OUT.mkdir(exist_ok=True)
    build_research().save(OUT / '选品调研对比表.xlsx')
    build_ops(a.price_list, a.sales).save(OUT / '上新排期与销售统计表.xlsx')
    print('products:', len(read_rows(a.price_list)), 'orders:', len(read_rows(a.sales)))
    print('wrote', OUT / '选品调研对比表.xlsx', OUT / '上新排期与销售统计表.xlsx')
