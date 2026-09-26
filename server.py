"""PetOps Studio — local-only 作图 + 上架 workbench. Python 3.11+, stdlib server, SQLite storage."""
from http.server import ThreadingHTTPServer, BaseHTTPRequestHandler
from pathlib import Path
from contextlib import contextmanager
import os, json, sqlite3, uuid, mimetypes, urllib.parse, base64, io, zipfile, threading, datetime as dt
import engine
import progress

ROOT = Path(__file__).resolve().parent
DATA = Path(os.environ.get('PETOPS_DATA', ROOT / 'data'))
DATA.mkdir(parents=True, exist_ok=True)
DB = DATA / 'petops.sqlite3'
PORT = 3310
HOSTS = [f'127.0.0.1:{PORT}', f'localhost:{PORT}']
MUTATION = threading.RLock()
JOBS = None


@contextmanager
def connect():
    c = sqlite3.connect(DB, timeout=20)
    c.row_factory = sqlite3.Row
    try:
        with c: yield c
    finally: c.close()


def init():
    with connect() as c:
        c.executescript('''
        PRAGMA journal_mode=WAL;
        CREATE TABLE IF NOT EXISTS products(id TEXT PRIMARY KEY, sku TEXT UNIQUE NOT NULL, body TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS jobs(id TEXT PRIMARY KEY, body TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS settings(id TEXT PRIMARY KEY, body TEXT NOT NULL);
        ''')


def records(table):
    with connect() as c:
        return [json.loads(r['body']) for r in c.execute(f'SELECT body FROM {table} ORDER BY rowid DESC')]


def save(table, obj):
    with connect() as c:
        if table == 'products':
            c.execute('INSERT INTO products VALUES(?,?,?) ON CONFLICT(id) DO UPDATE SET sku=excluded.sku,body=excluded.body', (obj['id'], obj['sku'], json.dumps(obj, ensure_ascii=False)))
        else:
            c.execute(f'INSERT INTO {table} VALUES(?,?) ON CONFLICT(id) DO UPDATE SET body=excluded.body', (obj['id'], json.dumps(obj, ensure_ascii=False)))


def config():
    return dict(engine.DEFAULTS, **next(iter(records('settings')), {}))


def get_product(pid):
    p = next((p for p in records('products') if p['id'] == pid), None)
    if not p: raise ValueError('商品不存在')
    return engine.migrate(p)


def file_path(rel):
    f = (DATA / rel).resolve()
    if not f.is_relative_to(DATA.resolve()) or not f.is_file(): raise ValueError('文件不存在')
    return f


def state():
    products = []
    for p in records('products'):
        p = engine.migrate(p)
        p['assets'] = engine.assets(p, DATA)
        p['boards'] = engine.boards(p, DATA)
        p['barcodeCheck'] = engine.barcode_status(p.get('ean'))
        products.append(p)
    jobs = records('jobs')
    for job in jobs: job['progress'] = progress.describe(job, DATA)
    return {'products': products, 'jobs': jobs[:200], 'settings': config(), 'capabilities': engine.capabilities(),
            'platforms': engine.SELLERS, 'statuses': engine.STATUSES, 'limits': engine.LIMITS}


def image_bytes(b64):
    raw = base64.b64decode(b64, validate=True)
    if len(raw) > 15 * 1024 * 1024: raise ValueError('图片最大 15MB')
    if raw.startswith(b'\x89PNG\r\n\x1a\n'): return raw, '.png'
    if raw.startswith(b'\xff\xd8\xff'): return raw, '.jpg'
    if raw[:4] == b'RIFF' and raw[8:12] == b'WEBP': return raw, '.webp'
    raise ValueError('仅支持 PNG、JPEG、WebP 图片')


# ---- Excel import: 上新排期与销售统计表.xlsx → sheet 上新排期 (also accepts the older 01_商品进价表 layout)
SCHEDULE_FIELDS = {'sku': 'sku', '商品ean码': 'ean', 'ean': 'ean', '商品简介': 'nameIt', '中文简介': 'nameZh', '税前进货价': 'cost', '仓库库存': 'stock',
                   '品类': 'category', '尺寸': 'dimensions', '重量': 'weight', '材质': 'material', '颜色': 'color', '卖点': 'sellingPoints', '品牌': 'brand'}
LISTING_FIELDS = {'商品id': ('AliExpress', 'productId'), '电商原价': ('AliExpress', 'price'), '日常促销价': ('AliExpress', 'promoPrice'),
                  'aliexpress': ('AliExpress', 'status'), 'amazon': ('Amazon', 'status'), 'tiktok': ('TikTok Shop', 'status')}
STATUS_WORDS = {'完成': '已上架', '进行中': '草稿', '未开始': '待上架', '下架': '已下架'}
UNKNOWN = {'', '待核实', '待补充', '-', '—', 'none'}


def norm(h):
    return str(h or '').replace('€', '').replace(' ', '').replace('　', '').strip().lower()


def cell_text(v):
    if v is None: return ''
    if isinstance(v, float) and v.is_integer(): v = int(v)
    if isinstance(v, (dt.date, dt.datetime)): return v.strftime('%Y-%m-%d')
    return str(v).strip()


def header_row(ws):
    for r in range(1, 12):
        values = [norm(c.value) for c in ws[r]]
        if 'sku' in values or '商品ean码' in values: return r, values
    raise ValueError(f'工作表“{ws.title}”找不到“商品EAN码”或“SKU”表头')


def import_workbook(raw):
    """Create or update products by SKU (EAN when the sheet has no SKU column). Only non-empty cells overwrite."""
    import openpyxl
    try: wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    except Exception: raise ValueError('无法读取 Excel 文件，请确认是 .xlsx 格式')
    sheet = next((wb[n] for n in ['上新排期', '排期'] if n in wb.sheetnames), wb.worksheets[0])
    r0, heads = header_row(sheet)
    note_col = heads.index('备注') if '备注' in heads else None
    existing = {p['sku']: engine.migrate(p) for p in records('products')}
    pending, skipped, created, listing_updates = {}, [], 0, 0
    for row in sheet.iter_rows(min_row=r0 + 1, values_only=True):
        cells = {h: cell_text(row[i]) for i, h in enumerate(heads) if i < len(row)}
        data = {SCHEDULE_FIELDS[h]: v for h, v in cells.items() if h in SCHEDULE_FIELDS and v.lower() not in UNKNOWN}
        sku = data.get('sku') or data.get('ean') or cells.get('商品id', '')
        if not sku: continue
        data['sku'] = sku
        if note_col is not None and '示例' in cells.get('备注', ''):
            skipped.append(f'{sku}：示例行'); continue
        old = pending.get(sku) or existing.get(sku)
        if not old and not (data.get('nameZh') or data.get('nameIt')):
            skipped.append(f'{sku}：缺少商品名称'); continue
        try: p = engine.product(data, old)
        except ValueError as e: skipped.append(f'{sku}：{e}'); continue
        changes = {}
        for h, (platform, key) in LISTING_FIELDS.items():
            value = cells.get(h, '')
            if key == 'status': value = STATUS_WORDS.get(value, '')
            if value: changes.setdefault(platform, {})[key] = value
        for platform, ch in changes.items():
            try: p = engine.listing_update(p, platform, ch); listing_updates += 1
            except ValueError as e: skipped.append(f'{sku} {platform}：{e}')
        if sku not in existing and sku not in pending: created += 1
        pending[sku] = p
    with connect() as c:
        for p in pending.values():
            c.execute('INSERT INTO products VALUES(?,?,?) ON CONFLICT(id) DO UPDATE SET sku=excluded.sku,body=excluded.body', (p['id'], p['sku'], json.dumps(p, ensure_ascii=False)))
    return {'created': created, 'updated': len(pending) - created, 'listingUpdates': listing_updates, 'skipped': skipped[:50]}


def mutate(path, d):
    if path == '/api/products':
        old = get_product(d['id']) if d.get('id') else None
        if old and str(d.get('sku', old['sku'])).strip() != old['sku'] and any(p['sku'] == str(d['sku']).strip() for p in records('products')): raise ValueError('SKU 已存在')
        p = engine.product(d, old); save('products', p); return p
    if path == '/api/delete-product':
        p = get_product(d['id'])
        if any(j['productId'] == p['id'] and j['status'] in ['queued', 'running'] for j in records('jobs')): raise ValueError('请先取消进行中的任务')
        with connect() as c: c.execute('DELETE FROM products WHERE id=?', (p['id'],))
        return {'ok': True}
    if path == '/api/reference':
        p = get_product(d['id'])
        if len(p['references']) >= 8: raise ValueError('每个商品最多 8 张参考图')
        raw, ext = image_bytes(d['data'])
        folder = DATA / 'references'; folder.mkdir(exist_ok=True)
        f = folder / (uuid.uuid4().hex + ext); f.write_bytes(raw)
        rel = str(f.relative_to(DATA)).replace('\\', '/')
        p['references'].append(rel); p['reference'] = p['references'][0]; save('products', p); return p
    if path == '/api/reference-order':
        p = get_product(d['id'])
        refs = [r for r in d.get('references', []) if r in p['references']]
        if set(refs) != set(p['references']) and not d.get('remove'): raise ValueError('参考图列表无效')
        p['references'] = refs; p['reference'] = refs[0] if refs else ''; save('products', p); return p
    if path == '/api/listing':
        p = engine.listing_update(get_product(d['id']), d['platform'], d.get('fields', {})); save('products', p); return p
    if path == '/api/review':
        p = get_product(d['id'])
        if not p.get('assetsReady'): raise ValueError('尚未生成完整图片')
        p['assetReview'] = bool(d['approved']); save('products', p); return p
    if path == '/api/adopt-edit':
        p = get_product(d['id'])
        j = next((j for j in records('jobs') if j['id'] == d['jobId'] and j['productId'] == p['id'] and j['kind'] == 'edit' and j['status'] == 'completed'), None)
        if not j or not j.get('edited'): raise ValueError('找不到可采用的修改图')
        file_path(j['edited'])
        source = next((a for a in engine.assets(p, DATA) if a['path'] == j['asset'] or a['originalPath'] == j['asset']), None)
        if not source: raise ValueError('原图已换版本，请基于当前图片重新修改')
        p.setdefault('assetOverrides', {})[source['originalPath']] = j['edited']
        p['assetReview'] = False; save('products', p); return p
    if path == '/api/revert-asset':
        p = get_product(d['id'])
        p.get('assetOverrides', {}).pop(d['originalPath'], None); p['assetReview'] = False; save('products', p); return p
    if path == '/api/jobs': return JOBS.submit(get_product(d['id']), d['kind'], d.get('extra', ''), d.get('asset', ''), d.get('platform', ''))
    if path == '/api/cancel-job': return JOBS.cancel(d['id'])
    if path == '/api/settings':
        s = engine.settings(d); save('settings', s); return s
    if path == '/api/import-xlsx':
        raw = base64.b64decode(d['data'], validate=True)
        if len(raw) > 20 * 1024 * 1024: raise ValueError('Excel 文件最大 20MB')
        return import_workbook(raw)
    if path == '/api/package':
        p = get_product(d['id'])
        folder = engine.write_listing_package(p, config(), DATA)
        return {'folder': str(folder)}
    raise ValueError('未知操作')


def bundle(p):
    folder = engine.write_listing_package(p, config(), DATA)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as z:
        for f in folder.rglob('*'):
            if f.is_file() and 'audit' not in f.parts: z.write(f, str(f.relative_to(folder)))
        for rel in engine.boards(p, DATA): z.write(file_path(rel), 'masters/' + Path(rel).name)
        z.writestr('README.txt', 'listing.json 与 ecom-listing-prep / ecom-aliexpress-publish / ecom-json-to-upload-xlsx 技能格式兼容。\n'
                                 'copy_*.md 为各平台意大利语文案；images/listing 为商品图（01 为白底主图），images/detail 为 AliExpress 详情图。\n'
                                 '上架前请核对 to_verify 中的待核实事实，以及各平台当前类目规则。')
    return buffer.getvalue()


class Handler(BaseHTTPRequestHandler):
    def log_message(self, fmt, *args):
        if '/api/state' not in (args[0] if args else ''): super().log_message(fmt, *args)

    def send(self, value, status=200, content_type='application/json; charset=utf-8', filename=None):
        payload = json.dumps(value, ensure_ascii=False).encode() if isinstance(value, (dict, list)) else value
        self.send_response(status)
        self.send_header('Content-Type', content_type)
        self.send_header('Content-Length', str(len(payload)))
        self.send_header('X-Content-Type-Options', 'nosniff')
        self.send_header('Cache-Control', 'no-store')
        if filename: self.send_header('Content-Disposition', "attachment; filename*=UTF-8''" + urllib.parse.quote(filename))
        self.end_headers()
        self.wfile.write(payload)

    def do_GET(self):
        if self.headers.get('Host') not in HOSTS: return self.send({'error': '仅允许本机访问'}, 403)
        try:
            parsed = urllib.parse.urlparse(self.path); path = urllib.parse.unquote(parsed.path); q = urllib.parse.parse_qs(parsed.query)
            if path == '/api/state': return self.send(state())
            if path == '/api/bundle':
                p = get_product(q['id'][0])
                return self.send(bundle(p), content_type='application/zip', filename=engine.safe_name(p['sku']) + '_上架资料包.zip')
            if path == '/api/backup':
                out = io.BytesIO()
                with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as z:
                    z.writestr('records.json', json.dumps(state(), ensure_ascii=False, indent=2))
                    for f in DATA.rglob('*'):
                        if f.is_file() and f.suffix in ['.png', '.jpg', '.webp', '.md', '.json', '.txt'] and 'tmp' not in f.parts: z.write(f, str(f.relative_to(DATA)))
                return self.send(out.getvalue(), content_type='application/zip', filename='PetOps_备份.zip')
            if path.startswith('/files/'):
                file = file_path(path[7:])
                if file.suffix.lower() not in ['.png', '.jpg', '.jpeg', '.webp', '.json', '.jsonl', '.txt', '.md']: raise ValueError('禁止下载此类型文件')
                image = file.suffix.lower() in ['.png', '.jpg', '.jpeg', '.webp']
                return self.send(file.read_bytes(), content_type=mimetypes.guess_type(file)[0] if image else 'text/plain; charset=utf-8',
                                 filename=file.name if 'download' in q else None)
            file = (ROOT / 'web' / ('index.html' if path == '/' else path.lstrip('/'))).resolve()
            if not file.is_relative_to(ROOT / 'web') or not file.is_file(): return self.send({'error': '不存在'}, 404)
            self.send(file.read_bytes(), content_type=mimetypes.guess_type(file)[0] or 'application/octet-stream')
        except (ValueError, KeyError) as e: self.send({'error': str(e)}, 400)

    def do_POST(self):
        origin = self.headers.get('Origin')
        if self.headers.get('Host') not in HOSTS or (origin and origin not in ['http://' + h for h in HOSTS]): return self.send({'error': '仅允许本机同源请求'}, 403)
        if not self.headers.get('Content-Type', '').startswith('application/json'): return self.send({'error': '需要 JSON 请求'}, 415)
        try:
            size = int(self.headers.get('Content-Length', '0'))
            if size < 1 or size > 30 * 1024 * 1024: raise ValueError('请求大小无效')
            data = json.loads(self.rfile.read(size))
            with MUTATION: result = mutate(self.path, data)
            self.send(result)
        except sqlite3.IntegrityError: self.send({'error': 'SKU 已存在，请编辑现有商品'}, 409)
        except (ValueError, KeyError, TypeError) as e: self.send({'error': str(e)}, 400)
        except Exception as e:
            print(type(e).__name__, str(e), flush=True); self.send({'error': '操作失败，请查看本机服务日志'}, 500)


if __name__ == '__main__':
    init()
    JOBS = engine.Jobs(DATA, records, save, config, MUTATION)
    print(f'PetOps Studio running at http://127.0.0.1:{PORT}', flush=True)
    ThreadingHTTPServer(('127.0.0.1', PORT), Handler).serve_forever()
